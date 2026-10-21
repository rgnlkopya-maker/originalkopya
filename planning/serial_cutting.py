from collections import defaultdict
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from datetime import date, timedelta

from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import OuterRef, Subquery
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from app_settings.access import has_feature_access, order_scope_allows
from core.models import Order, OrderEvent
from core.order_list_enhanced import STAGE_TRANSLATIONS, _normalize, _transfer_status
from .models import SerialCutPlan


def _can_edit(user):
    return has_feature_access(user, "planning.edit_calendar")


@login_required
def serial_cutting_page(request):
    raw = request.GET.get("date")
    try:
        selected = date.fromisoformat(raw) if raw else date.today()
    except ValueError:
        selected = date.today()
    monday = selected - timedelta(days=selected.weekday())
    weekdays = [monday + timedelta(days=i) for i in range(5)]
    plans = SerialCutPlan.objects.filter(planned_date__range=(monday, weekdays[-1])).select_related("order")
    grouped = {day.isoformat(): {} for day in weekdays}
    for plan in plans:
        order = plan.order
        if not order_scope_allows(request.user, order):
            continue
        key = (order.urun_kodu or "", order.renk or "", order.beden or "")
        bucket = grouped[plan.planned_date.isoformat()].setdefault(key, {
            "product": key[0], "color": key[1], "size": key[2], "quantity": 0, "order_ids": [], "plans": []
        })
        bucket["quantity"] += plan.quantity
        bucket["order_ids"].append(order.id)
        bucket["plans"].append(plan)
    columns = []
    week_total = 0
    for day in weekdays:
        rows = sorted(grouped[day.isoformat()].values(), key=lambda r: (r["product"], r["color"], r["size"]))
        total = sum(row["quantity"] for row in rows)
        week_total += total
        products = defaultdict(int)
        for row in rows:
            products[row["product"]] += row["quantity"]
        columns.append({"date": day, "rows": rows, "total": total, "products": sorted(products.items())})
    planned_ids = set(SerialCutPlan.objects.values_list("order_id", flat=True))
    last_event = OrderEvent.objects.filter(order=OuterRef("pk")).exclude(event_type="order_update").exclude(stage__in=["satis_fiyati", "ekstra_maliyet", "maliyet_indirimi", "maliyet_override", "maliyet_uygulanan"]).order_by("-timestamp", "-id")
    orders = Order.objects.filter(is_active=True).exclude(siparis_tipi="MALZEME").exclude(id__in=planned_ids).select_related("musteri").annotate(status_stage=Subquery(last_event.values("stage")[:1]), status_value=Subquery(last_event.values("value")[:1]), status_parca=Subquery(last_event.values("parca")[:1])).order_by("-id")[:1500]
    orders = [order for order in orders if order_scope_allows(request.user, order)]
    for order in orders:
        stage, value = order.status_stage, order.status_value
        if not stage or not value:
            order.cut_status = "-"
        else:
            normalized = (_normalize(stage), _normalize(value))
            fallback = stage.replace("_durum", "").replace("_", " ").title() + " → " + value.replace("_", " ").title()
            order.cut_status = _transfer_status(stage, value, order.status_parca) or STAGE_TRANSLATIONS.get(normalized) or fallback
    return render(request, "planning/serial_cutting.html", {
        "columns": columns, "week_total": week_total, "orders": orders,
        "can_edit": _can_edit(request.user),
        "prev_date": (monday-timedelta(days=7)).isoformat(),
        "next_date": (monday+timedelta(days=7)).isoformat(),
    })


@login_required
def serial_cutting_weekly_excel(request):
    """One printable color-by-size matrix per product for the selected workweek."""
    raw = request.GET.get("date")
    try:
        selected = date.fromisoformat(raw) if raw else date.today()
    except ValueError:
        selected = date.today()
    monday = selected - timedelta(days=selected.weekday())
    friday = monday + timedelta(days=4)
    plans = SerialCutPlan.objects.filter(
        planned_date__range=(monday, friday)
    ).select_related("order")
    grouped = defaultdict(lambda: defaultdict(int))
    for plan in plans:
        if not order_scope_allows(request.user, plan.order):
            continue
        product = (plan.order.urun_kodu or "Kod belirtilmemiş").strip()
        color = (plan.order.renk or "Renk belirtilmemiş").strip()
        size = str(plan.order.beden or "Beden belirtilmemiş").strip()
        grouped[product][(color, size)] += plan.quantity

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()
    for product in sorted(grouped, key=str.casefold):
        # Excel sheet titles must be unique, <=31 chars and exclude special characters.
        title_base = "".join("_" if ch in "[]:*?/\\\\ " else ch for ch in product).strip("_")[:31] or "Urun"
        title = title_base
        suffix = 2
        while title.casefold() in used_titles:
            ending = "_" + str(suffix)
            title = title_base[:31-len(ending)] + ending
            suffix += 1
        used_titles.add(title.casefold())
        sheet = workbook.create_sheet(title)
        matrix = grouped[product]
        colors = sorted({color for color, _ in matrix}, key=str.casefold)
        sizes = sorted({size for _, size in matrix}, key=lambda s: (not s.isdigit(), int(s) if s.isdigit() else s.casefold()))
        last_col = len(sizes) + 2
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        sheet.cell(1, 1, product)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        sheet.cell(2, 1, monday.strftime("%d.%m.%Y") + " - " + friday.strftime("%d.%m.%Y"))
        headers = ["RENK / BEDEN"] + sizes + ["TOPLAM"]
        for col, label in enumerate(headers, 1):
            sheet.cell(4, col, label)
        for row_idx, color in enumerate(colors, 5):
            sheet.cell(row_idx, 1, color)
            for col_idx, size in enumerate(sizes, 2):
                count = matrix.get((color, size), 0)
                sheet.cell(row_idx, col_idx, count if count else None)
            sheet.cell(row_idx, last_col, sum(matrix.get((color, size), 0) for size in sizes))
        total_row = 5 + len(colors)
        sheet.cell(total_row, 1, "TOPLAM")
        for col_idx, size in enumerate(sizes, 2):
            sheet.cell(total_row, col_idx, sum(matrix.get((color, size), 0) for color in colors))
        sheet.cell(total_row, last_col, sum(matrix.values()))

        thin = Side(style="thin", color="333333")
        for row in sheet.iter_rows(min_row=1, max_row=total_row, max_col=last_col):
            for cell in row:
                cell.alignment = Alignment(horizontal="center" if cell.column > 1 else "left", vertical="center")
                cell.font = Font(name="Calibri", size=11, color="000000",
                                 bold=cell.row in (1, 4, total_row))
                if cell.row >= 4:
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet.cell(1, 1).font = Font(name="Calibri", size=16, bold=True, color="000000")
        sheet.cell(2, 1).font = Font(name="Calibri", size=10, color="555555")
        sheet.column_dimensions["A"].width = 23
        for col in range(2, last_col + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 12
        sheet.row_dimensions[1].height = 27
        sheet.row_dimensions[4].height = 23
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B5"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:{get_column_letter(last_col)}{total_row}"
    if not workbook.sheetnames:
        sheet = workbook.create_sheet("Bos Plan")
        sheet["A1"] = "Bu hafta için kesim planı bulunamadı."
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="seri-kesim-{monday.isoformat()}.xlsx"'
    return response


@login_required
@require_POST
def serial_cutting_add(request):
    if not _can_edit(request.user):
        return HttpResponseForbidden("Planlama yetkiniz yok.")
    try:
        day = date.fromisoformat(request.POST.get("planned_date", ""))
    except (ValueError, TypeError):
        return HttpResponseForbidden("Geçersiz tarih.")
    if day.weekday() > 4:
        return HttpResponseForbidden("Yalnızca hafta içi planlanabilir.")
    raw_ids = request.POST.getlist("order_ids")
    if not raw_ids or len(raw_ids) > 1500:
        return HttpResponseForbidden("Sipariş seçiniz.")
    try:
        ids = list(dict.fromkeys(int(value) for value in raw_ids))
    except (ValueError, TypeError):
        return HttpResponseForbidden("Geçersiz sipariş seçimi.")
    with transaction.atomic():
        orders = list(Order.objects.select_for_update().filter(id__in=ids, is_active=True).exclude(siparis_tipi="MALZEME"))
        if len(orders) != len(ids) or any(not order_scope_allows(request.user, order) for order in orders):
            return HttpResponseForbidden("Siparişlerden birine erişim yok.")
        if SerialCutPlan.objects.filter(order_id__in=ids).exists():
            return HttpResponseForbidden("Seçilen siparişlerden biri zaten kesim planında.")
        SerialCutPlan.objects.bulk_create([
            SerialCutPlan(order=order, planned_date=day, quantity=1, created_by=request.user)
            for order in orders
        ])
    return redirect(f"{reverse('serial_cutting_page')}?date={day.isoformat()}")

@login_required
@require_POST
def serial_cutting_remove(request, plan_id):
    if not _can_edit(request.user):
        return HttpResponseForbidden("Planlama yetkiniz yok.")
    plan = get_object_or_404(SerialCutPlan, pk=plan_id)
    if not order_scope_allows(request.user, plan.order):
        return HttpResponseForbidden("Sipariş yetkiniz yok.")
    day = plan.planned_date
    plan.delete()
    return redirect(f"{reverse('serial_cutting_page')}?date={day.isoformat()}")


@login_required
def serial_cutting_statuses(request):
    """Fetch current order-list statuses only when a cutting detail dialog opens."""
    raw_ids = request.GET.getlist("order_id")
    if not raw_ids or len(raw_ids) > 150:
        return JsonResponse({"error": "Geçersiz sipariş seçimi."}, status=400)
    try:
        ids = list(dict.fromkeys(int(value) for value in raw_ids))
    except (ValueError, TypeError):
        return JsonResponse({"error": "Geçersiz sipariş seçimi."}, status=400)
    plans = SerialCutPlan.objects.filter(order_id__in=ids).select_related("order")
    visible_ids = {plan.order_id for plan in plans if order_scope_allows(request.user, plan.order)}
    if len(visible_ids) != len(ids):
        return JsonResponse({"error": "Siparişlere erişim yok."}, status=403)
    events = (OrderEvent.objects.filter(order_id__in=visible_ids)
              .exclude(event_type="order_update")
              .exclude(stage__in=["satis_fiyati", "ekstra_maliyet", "maliyet_indirimi", "maliyet_override", "maliyet_uygulanan"])
              .order_by("order_id", "-timestamp", "-id"))
    latest = {}
    for event in events:
        if event.order_id not in latest:
            latest[event.order_id] = event
    statuses = {}
    for order_id in visible_ids:
        event = latest.get(order_id)
        if not event or not event.stage or not event.value:
            statuses[str(order_id)] = "-"
            continue
        transfer = _transfer_status(event.stage, event.value, event.parca)
        key = (_normalize(event.stage), _normalize(event.value))
        if transfer:
            status = transfer
        elif key in STAGE_TRANSLATIONS:
            status = STAGE_TRANSLATIONS[key]
        else:
            stage = event.stage.replace("_durum", "").replace("_fason_durumu", " Fason").replace("_", " ").title()
            status = f"{stage} → {event.value.replace('_', ' ').title()}"
        statuses[str(order_id)] = status
    response = JsonResponse({"statuses": statuses})
    response["Cache-Control"] = "no-store"
    return response
