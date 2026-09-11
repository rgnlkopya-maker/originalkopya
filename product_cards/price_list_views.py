from decimal import Decimal, InvalidOperation
from io import BytesIO
import urllib.request
import xml.etree.ElementTree as ET

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import ExchangeRate, PriceListSettings, ProductCard


def _can_manage(user):
    return user.is_superuser or user.groups.filter(name__in=["patron", "mudur"]).exists()


def fetch_price_list_tcmb_rates():
    request = urllib.request.Request("https://www.tcmb.gov.tr/kurlar/today.xml", headers={"User-Agent": "MoliApp/1.0"})
    with urllib.request.urlopen(request, timeout=12) as response:
        root = ET.fromstring(response.read())
    values = {}
    for code in ("USD", "EUR"):
        node = root.find(f".//Currency[@CurrencyCode='{code}']")
        text = (node.findtext("ForexSelling") or node.findtext("BanknoteSelling") if node is not None else None)
        if not text:
            raise RuntimeError(f"TCMB {code} satış kuru alınamadı.")
        values[code] = Decimal(text.strip().replace(",", "."))
    source_date = root.attrib.get("Date", "")
    today = timezone.localdate()
    rate, _ = ExchangeRate.objects.update_or_create(rate_date=today, defaults={"usd_try": values["USD"], "eur_try": values["EUR"], "source_date": source_date})
    settings = PriceListSettings.get_solo()
    settings.usd_try = values["USD"]; settings.eur_try = values["EUR"]; settings.rate_source = "TCMB"; settings.rate_source_date = source_date; settings.rate_checked_at = timezone.now()
    settings.save(update_fields=["usd_try", "eur_try", "rate_source", "rate_source_date", "rate_checked_at", "updated_at"])
    return rate


def _ensure_price_rates(settings):
    checked = settings.rate_checked_at
    if checked and timezone.localdate(checked) == timezone.localdate() and settings.usd_try > 1 and settings.eur_try > 1:
        return None
    try:
        fetch_price_list_tcmb_rates(); settings.refresh_from_db(); return None
    except Exception as exc:
        latest = ExchangeRate.objects.order_by("-rate_date", "-fetched_at").first()
        if latest and settings.usd_try <= 1:
            settings.usd_try = latest.usd_try; settings.eur_try = latest.eur_try; settings.rate_source = "Son geçerli TCMB"; settings.rate_source_date = latest.source_date; settings.save()
        return str(exc)


def _real_profit_rate(profit_rate, discount_rate):
    multiplier = (Decimal("1") + profit_rate / Decimal("100")) * (Decimal("1") - discount_rate / Decimal("100"))
    return ((multiplier - Decimal("1")) * Decimal("100")).quantize(Decimal("0.01"))


def _price_rows(settings, active=True):
    profit = settings.profit_rate / Decimal("100"); discount = settings.discount_rate / Decimal("100"); monthly = settings.monthly_term_rate / Decimal("100"); usd_rate = settings.usd_try; eur_rate = settings.eur_try
    cards = ProductCard.objects.select_related("urun").prefetch_related("materials__material").filter(price_list_active=active).order_by("urun__kod")
    rows = []
    for card in cards:
        cost = card.toplam_maliyet; with_profit = cost * (Decimal("1") + profit); discounted = with_profit * (Decimal("1") - discount)
        rows.append({"id": card.id, "code": card.urun.kod, "cash": discounted.quantize(Decimal("0.01")), "term3": (discounted * (Decimal("1") + monthly * 3)).quantize(Decimal("0.01")), "term6": (discounted * (Decimal("1") + monthly * 6)).quantize(Decimal("0.01")), "term9": (discounted * (Decimal("1") + monthly * 9)).quantize(Decimal("0.01")), "usd": (discounted / usd_rate).quantize(Decimal("0.01")) if usd_rate > 0 else None, "eur": (discounted / eur_rate).quantize(Decimal("0.01")) if eur_rate > 0 else None})
    return rows


@login_required
def price_list(request):
    if not _can_manage(request.user): return HttpResponseForbidden("Bu sayfaya erişim yetkiniz yok.")
    settings = PriceListSettings.get_solo()
    if settings.profit_rate <= 0 and settings.discount_rate > 0:
        settings.discount_rate = Decimal("0"); settings.save(update_fields=["discount_rate", "updated_at"])
    rate_error = _ensure_price_rates(settings); show_inactive = request.GET.get("durum") == "pasif"
    context = {"settings": settings, "rows": _price_rows(settings, active=not show_inactive), "show_inactive": show_inactive, "inactive_count": ProductCard.objects.filter(price_list_active=False).count(), "rate_error": rate_error, "real_profit_rate": _real_profit_rate(settings.profit_rate, settings.discount_rate), "showroom_mode": bool(getattr(request, "_showroom_mode", False))}
    if context["showroom_mode"]:
        from core.models import Beden, Musteri, Renk, URUN_TIPI_CHOICES, UrunKod
        musteriler = Musteri.objects.filter(aktif=True).order_by("ad"); renkler = Renk.objects.filter(aktif=True).order_by("ad"); bedenler = Beden.objects.filter(aktif=True).order_by("ad"); urun_kodlari = UrunKod.objects.filter(aktif=True).order_by("kod")
        context.update({"musteriler": musteriler, "renkler": renkler, "bedenler": bedenler, "urun_kodlari": urun_kodlari, "aktif_musteriler": musteriler, "aktif_renkler": renkler, "aktif_bedenler": bedenler, "aktif_urun_kodlari": urun_kodlari, "urun_tipi_secenekleri": URUN_TIPI_CHOICES})
    return render(request, "product_cards/price_list.html", context)


@login_required
@require_POST
def save_price_list_settings(request):
    if not _can_manage(request.user): return JsonResponse({"ok": False, "message": "Yetkiniz yok."}, status=403)
    settings = PriceListSettings.get_solo(); fields = {"profit_rate": "Kâr oranı", "discount_rate": "İndirim oranı", "monthly_term_rate": "Aylık vade farkı", "usd_try": "USD kuru", "eur_try": "EUR kuru"}; parsed = {}
    try:
        for field, label in fields.items():
            raw = (request.POST.get(field) or "").strip().replace(",", ".")
            if not raw and field in {"profit_rate", "discount_rate", "monthly_term_rate"}: raw = "0"
            value = Decimal(raw)
            if value < 0: raise ValueError(f"{label} negatif olamaz.")
            if field == "discount_rate" and value > 100: raise ValueError("İndirim oranı %100'den büyük olamaz.")
            if field in {"usd_try", "eur_try"} and value <= 0: raise ValueError(f"{label} sıfır olamaz.")
            parsed[field] = value
    except (InvalidOperation, ValueError) as exc: return JsonResponse({"ok": False, "message": str(exc)}, status=400)
    if parsed["discount_rate"] > 0 and parsed["profit_rate"] <= 0: return JsonResponse({"ok": False, "message": "İndirim uygulamak için önce kâr oranı girmelisiniz."}, status=400)
    for field, value in parsed.items(): setattr(settings, field, value)
    if request.POST.get("changed_field") in {"usd_try", "eur_try"}: settings.rate_source = "Manuel senaryo"; settings.rate_checked_at = timezone.now()
    settings.save()
    return JsonResponse({"ok": True, "message": "Otomatik kaydedildi", "source": settings.rate_source, "checked_at": timezone.localtime(settings.rate_checked_at).strftime("%d.%m.%Y %H:%M") if settings.rate_checked_at else "", "real_profit_rate": str(_real_profit_rate(settings.profit_rate, settings.discount_rate)), "rows": [{key: str(value) if isinstance(value, Decimal) else value for key, value in row.items()} for row in _price_rows(settings)]})


@login_required
@require_POST
def toggle_price_list_status(request):
    if not _can_manage(request.user): return HttpResponseForbidden("Bu işlem için yetkiniz yok.")
    card = get_object_or_404(ProductCard, pk=request.POST.get("card_id")); card.price_list_active = request.POST.get("active") == "1"; card.save(update_fields=["price_list_active", "updated_at"])
    if request.POST.get("return_status") == "pasif": return redirect(f'{reverse("price_list")}?durum=pasif')
    return redirect("price_list")


@login_required
def export_price_list_excel(request):
    if not _can_manage(request.user): return HttpResponseForbidden("Bu işlem için yetkiniz yok.")
    settings = PriceListSettings.get_solo(); rows = _price_rows(settings, active=True); workbook = Workbook(); sheet = workbook.active; sheet.title = "Fiyat Listesi"; headers = ["Ürün Kodu", "Nakit TL", "3 Ay TL", "6 Ay TL", "9 Ay TL", "USD", "EUR"]; sheet.append(headers)
    for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="4A63FF"); cell.alignment = Alignment(horizontal="center")
    for row in rows: sheet.append([row["code"], row["cash"], row["term3"], row["term6"], row["term9"], row["usd"], row["eur"]])
    for column, width in {"A": 22, "B": 16, "C": 16, "D": 16, "E": 16, "F": 14, "G": 14}.items(): sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows(min_row=2, min_col=2, max_col=7):
        for cell in cells: cell.number_format = '#,##0.00'
    sheet.freeze_panes = "A2"; output = BytesIO(); workbook.save(output); response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); response["Content-Disposition"] = 'attachment; filename="fiyat-listesi.xlsx"'; return response
